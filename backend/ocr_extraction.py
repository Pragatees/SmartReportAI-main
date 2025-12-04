import fitz  # PyMuPDF
import pytesseract
from PIL import Image
import io
import os
from fastapi import APIRouter, UploadFile, File
from fastapi.responses import JSONResponse
from docx import Document

# ✅ Set Tesseract executable path
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

router = APIRouter(prefix="/ocr", tags=["OCR"])


@router.post("/extract-text")
async def extract_text(file: UploadFile = File(...)):
    try:
        contents = await file.read()
        file_extension = os.path.splitext(file.filename)[1].lower()
        text = ""

        if file_extension == ".pdf":
            doc = fitz.open(stream=contents, filetype="pdf")
            for page_num in range(doc.page_count):
                page = doc[page_num]
                page_text = page.get_text()
                if page_text.strip():
                    text += page_text + "\n"
                else:
                    img_list = page.get_images(full=True)
                    for img_index, img in enumerate(img_list):
                        xref = img[0]
                        base_image = doc.extract_image(xref)
                        image_bytes = base_image["image"]
                        try:
                            img_obj = Image.open(io.BytesIO(image_bytes))
                            img_text = pytesseract.image_to_string(img_obj)
                            text += img_text + "\n"
                        except Exception as img_e:
                            print(f"Image OCR error on page {page_num}: {img_e}")
                text += f"\n--- End of Page {page_num + 1} ---\n"
            doc.close()

        elif file_extension in [".png", ".jpg", ".jpeg", ".tif", ".tiff"]:
            img = Image.open(io.BytesIO(contents))
            text = pytesseract.image_to_string(img)

        elif file_extension in [".doc", ".docx"]:
            try:
                doc = Document(io.BytesIO(contents))
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip():
                        text += paragraph.text + "\n"
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                text += cell.text + "\n"
            except Exception as doc_e:
                print(f"Word document processing error: {doc_e}")
                return JSONResponse(content={"error": f"Failed to process Word document: {str(doc_e)}"}, status_code=400)

        else:
            return JSONResponse(content={"error": f"Unsupported file format: {file_extension}"}, status_code=400)

        return {"filename": file.filename, "text": text}

    except Exception as e:
        print(f"General error: {str(e)}")
        return JSONResponse(content={"error": str(e)}, status_code=500)